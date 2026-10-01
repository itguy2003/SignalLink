"""기관 엑셀·CSV 일괄 등록 파서.

요구사항: "기관등록은 엑셀파일로 일괄 등록할 수 있어야 한다."
관리자 화면의 학교 등록은 CSV 도 받는다 — 나이스 등에서 내려받은 목록이
흔히 CSV 라서다.

첫 행을 헤더로 보고 컬럼명을 매핑한다. 컬럼 순서는 자유이며,
알 수 없는 컬럼은 무시한다. 필수는 기관명 하나뿐이다.
"""

from __future__ import annotations

import io
from typing import Any

from openpyxl import load_workbook

from app.models.enums import OrgType, SchoolLevel

#: 엑셀 헤더 -> 모델 필드
COLUMN_MAP: dict[str, str] = {
    "기관명": "name",
    # 위경도 자료처럼 다른 출처의 파일은 흔히 '학교명' 이라고 적는다.
    "학교명": "name",
    "주소": "address",
    "기관타입": "org_type",
    "학교단위": "school_level",
    "위도": "latitude",
    "경도": "longitude",
    "시도교육청코드": "edu_office_code",
    "학교코드": "school_code",
    # 나이스 학교기본정보 배포 파일의 표준 헤더들.
    "시도교육청명": "edu_office_name",
    "행정코드": "school_code",
    "행정표준코드": "school_code",
    "학교종류명": "school_level",
    "학교 종류명": "school_level",
    "시도명": "sido_name",
    "관할조직명": "jurisdiction_name",
    "도로명주소": "address",
    "도로명상세주소": "address_detail",
    "전화번호": "phone",
    "홈페이지주소": "homepage",
}

TEMPLATE_HEADERS = list(COLUMN_MAP)


def _normalize_school_level(value: str | None, valid_levels: set[str]) -> str | None:
    """'학교종류명'(초등학교/중학교/고등학교/특수학교 등)을 학교급으로.

    이미 학교급이거나 못 알아보는 값은 그대로 둔다 — 검증이 잡는다.
    """
    if value is None or value in valid_levels:
        return value
    if "초등" in value:
        return "초"
    if "고등" in value:
        return "고"
    if "중" in value:
        return "중"
    if value.endswith("학교"):
        return "기타"
    return value


def _as_float(value: Any) -> float | None:
    if value in (None, ""):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        raise ValueError(f"숫자가 아닙니다: {value!r}") from None


def _as_str(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def parse_organization_workbook(
    content: bytes,
    default_org_type: str | None = None,
    allowed_types: set[str] | None = None,
) -> tuple[list[dict[str, Any]], list[str]]:
    """엑셀 바이트를 읽어 (등록할 행 목록, 오류 메시지 목록)을 반환한다."""
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - openpyxl 은 다양한 예외를 던진다
        return [], [f"엑셀 파일을 읽을 수 없습니다: {exc}"]

    sheet = workbook.active
    if sheet is None:
        return [], ["시트를 찾을 수 없습니다."]

    rows_iter = sheet.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        return [], ["빈 파일입니다."]

    header_index: dict[int, str] = {}
    for index, cell in enumerate(header_row):
        key = _as_str(cell)
        if key and key in COLUMN_MAP:
            header_index[index] = COLUMN_MAP[key]

    if "name" not in header_index.values():
        return [], [f"'기관명' 컬럼이 없습니다. 필요한 헤더: {', '.join(TEMPLATE_HEADERS)}"]

    valid_types = allowed_types or {str(t) for t in OrgType}
    # '기타'는 특수학교처럼 학교급 셋에 안 드는 학교의 자리다(.NET 파서와 동일).
    valid_levels = {str(level) for level in SchoolLevel} | {"기타"}

    parsed: list[dict[str, Any]] = []
    errors: list[str] = []

    for row_number, row in enumerate(rows_iter, start=2):
        if row is None or all(cell in (None, "") for cell in row):
            continue

        record: dict[str, Any] = {}
        try:
            for index, field in header_index.items():
                value = row[index] if index < len(row) else None
                if field in ("latitude", "longitude"):
                    record[field] = _as_float(value)
                else:
                    record[field] = _as_str(value)

            if not record.get("name"):
                raise ValueError("기관명이 비어 있습니다.")

            org_type = record.get("org_type") or default_org_type or str(OrgType.ORG)
            if org_type not in valid_types:
                allowed = ", ".join(sorted(valid_types))
                raise ValueError(f"기관타입이 올바르지 않습니다: {org_type} (허용: {allowed})")
            record["org_type"] = org_type

            level = _normalize_school_level(record.get("school_level"), valid_levels)
            record["school_level"] = level
            if level and level not in valid_levels:
                allowed = ", ".join(sorted(valid_levels))
                raise ValueError(f"학교단위가 올바르지 않습니다: {level} (허용: {allowed})")

            parsed.append(record)
        except ValueError as exc:
            errors.append(f"{row_number}행: {exc}")

    workbook.close()
    return parsed, errors


def _decode_csv(content: bytes) -> str:
    """CSV 바이트를 문자열로. 한국 환경의 두 인코딩을 순서대로 시도한다.

    엑셀이 저장한 CSV 는 cp949 인 경우가 많고, 요즘 도구는 UTF-8(BOM 포함)을
    쓴다. BOM 이 있으면 utf-8-sig 가 벗겨 준다.
    """
    for encoding in ("utf-8-sig", "cp949"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("파일 인코딩을 해석할 수 없습니다. UTF-8 또는 CP949 로 저장하세요.")


def parse_organization_csv(
    content: bytes,
    default_org_type: str | None = None,
    allowed_types: set[str] | None = None,
) -> tuple[list[dict[str, Any]], list[str]]:
    """CSV 바이트를 읽어 (등록할 행 목록, 오류 메시지 목록)을 반환한다.

    헤더 규칙은 엑셀과 같다(``COLUMN_MAP``). 형식만 다르고 규칙이 갈리면
    "엑셀로는 되는데 CSV 로는 안 되는" 헤더가 생긴다.
    """
    import csv as _csv

    try:
        text = _decode_csv(content)
    except ValueError as exc:
        return [], [str(exc)]

    rows_iter = _csv.reader(io.StringIO(text))
    try:
        header_row = next(rows_iter)
    except StopIteration:
        return [], ["빈 파일입니다."]

    header_index: dict[int, str] = {}
    for index, cell in enumerate(header_row):
        key = _as_str(cell)
        if key and key in COLUMN_MAP:
            header_index[index] = COLUMN_MAP[key]

    if "name" not in header_index.values():
        return [], [f"'기관명' 컬럼이 없습니다. 필요한 헤더: {', '.join(TEMPLATE_HEADERS)}"]

    valid_types = allowed_types or {str(t) for t in OrgType}
    # '기타'는 특수학교처럼 학교급 셋에 안 드는 학교의 자리다(.NET 파서와 동일).
    valid_levels = {str(level) for level in SchoolLevel} | {"기타"}

    parsed: list[dict[str, Any]] = []
    errors: list[str] = []

    for row_number, row in enumerate(rows_iter, start=2):
        if not row or all(_as_str(cell) is None for cell in row):
            continue

        record: dict[str, Any] = {}
        try:
            for index, field in header_index.items():
                value = row[index] if index < len(row) else None
                if field in ("latitude", "longitude"):
                    record[field] = _as_float(value)
                else:
                    record[field] = _as_str(value)

            if not record.get("name"):
                raise ValueError("기관명이 비어 있습니다.")

            org_type = record.get("org_type") or default_org_type or str(OrgType.ORG)
            if org_type not in valid_types:
                allowed = ", ".join(sorted(valid_types))
                raise ValueError(f"기관타입이 올바르지 않습니다: {org_type} (허용: {allowed})")
            record["org_type"] = org_type

            level = _normalize_school_level(record.get("school_level"), valid_levels)
            record["school_level"] = level
            if level and level not in valid_levels:
                allowed = ", ".join(sorted(valid_levels))
                raise ValueError(f"학교단위가 올바르지 않습니다: {level} (허용: {allowed})")

            parsed.append(record)
        except ValueError as exc:
            errors.append(f"{row_number}행: {exc}")

    return parsed, errors


#: 보조 파일에서 본 파일로 가져오는 필드 — 두 번째 파일의 몫은 좌표뿐이다.
SUPPLEMENT_FIELDS = ("latitude", "longitude")


def merge_supplement(
    base_rows: list[dict[str, Any]],
    extra_rows: list[dict[str, Any]],
) -> list[str]:
    """보조 파일의 **위도·경도만** 본 파일 행에 채워 넣는다.

    학교 정보와 좌표가 서로 다른 파일로 오는 경우다(요구사항: 복수 출처 결합).
    짝은 (시도교육청코드, 학교코드)가 양쪽에 있으면 그것으로, 아니면 이름으로
    맞춘다 — 코드가 더 정확하지만 좌표 파일에는 이름만 있는 경우가 많다.

    본 파일이 이미 가진 값은 덮지 않는다 — 본 파일이 정본이다.
    짝을 못 찾은 보조 행은 오류 목록으로 돌려준다. 조용히 버리면 좌표가
    빠진 학교를 아무도 모른다.
    """
    by_code: dict[tuple[str, str], dict[str, Any]] = {}
    by_name: dict[str, dict[str, Any]] = {}
    for row in base_rows:
        if row.get("edu_office_code") and row.get("school_code"):
            by_code[(row["edu_office_code"], row["school_code"])] = row
        if row.get("name"):
            by_name[row["name"]] = row

    errors: list[str] = []
    for at, extra in enumerate(extra_rows, start=2):
        target = None
        if extra.get("edu_office_code") and extra.get("school_code"):
            target = by_code.get((extra["edu_office_code"], extra["school_code"]))
        if target is None and extra.get("name"):
            target = by_name.get(extra["name"])
        if target is None:
            label = extra.get("name") or extra.get("school_code") or f"{at}행"
            errors.append(f"보조 파일: '{label}' 과 짝이 되는 행을 본 파일에서 찾지 못했습니다.")
            continue
        for field in SUPPLEMENT_FIELDS:
            value = extra.get(field)
            if value is not None and target.get(field) in (None, ""):
                target[field] = value
    return errors
